"""Build the downloadable Opening Stock XLSX template the importer accepts."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from common.i18n import t
from dlux.translations import get_current_language_code

from .models import Category, Product
from .opening_stock_import import MAX_IMPORT_ROWS, TEMPLATE_SHEET, column_guide

LAST_ROW = MAX_IMPORT_ROWS + 1
REQUIRED_FILL = PatternFill("solid", fgColor="1F6F43")
OPTIONAL_FILL = PatternFill("solid", fgColor="D9E2EC")
WRAP = Alignment(wrap_text=True, vertical="top")
COLUMN_WIDTHS = {"name": 32, "alias": 26, "category": 30, "barcode": 18, "size": 16}
DECIMAL_KEYS = {"cost_usd", "markup_percent", "price_usd", "price_lyd_override", "quantity"}


def _list_validation(sheet, column, formula, label):
    validation = DataValidation(
        type="list", formula1=formula, allow_blank=True,
        showErrorMessage=True, errorStyle="stop",
        errorTitle=label[:32], error=t("ui_excel_pick_from_list", "Choose a value from the list."),
    )
    validation.add(f"{column}2:{column}{LAST_ROW}")
    sheet.add_data_validation(validation)


def _write_list(sheet, column, title, values):
    sheet[f"{column}1"] = title
    sheet[f"{column}1"].font = Font(bold=True)
    for row, value in enumerate(values, start=2):
        sheet[f"{column}{row}"] = value
    return f"=Lists!${column}$2:${column}${max(len(values), 1) + 1}"


def build_template():
    guide = column_guide()
    workbook = Workbook()

    data = workbook.active
    data.title = TEMPLATE_SHEET
    instructions = workbook.create_sheet("Instructions")
    lists = workbook.create_sheet("Lists")
    lists.sheet_state = "hidden"

    unit_ref = _write_list(lists, "A", "Unit", [label for _code, label in Product.UNIT_CHOICES])
    color_ref = _write_list(lists, "B", "Color", [label for _code, label in Product.COLOR_CHOICES])
    categories = sorted(str(category) for category in Category.objects.filter(is_service=False).select_related("parent"))
    category_ref = _write_list(lists, "C", "Category", categories) if categories else None

    for index, column in enumerate(guide, start=1):
        letter = get_column_letter(index)
        cell = data.cell(row=1, column=index, value=column["label"])
        cell.font = Font(bold=True, color="FFFFFF" if column["required"] else "1F2933")
        cell.fill = REQUIRED_FILL if column["required"] else OPTIONAL_FILL
        status = t("ui_excel_required", "Required") if column["required"] else t("ui_excel_optional", "Optional")
        note = Comment(f"{status}. {column['description']}", "Sales CRM")
        note.width, note.height = 320, 140
        cell.comment = note
        data.column_dimensions[letter].width = COLUMN_WIDTHS.get(column["key"], 15)

        key = column["key"]
        if key == "unit":
            _list_validation(data, letter, unit_ref, column["label"])
        elif key == "color":
            _list_validation(data, letter, color_ref, column["label"])
        elif key == "currency":
            _list_validation(data, letter, '"USD,EUR"', column["label"])
        elif key == "category" and category_ref:
            _list_validation(data, letter, category_ref, column["label"])
        elif key in DECIMAL_KEYS:
            validation = DataValidation(
                type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                showErrorMessage=True, errorStyle="stop", errorTitle=column["label"][:32],
                error=t("ui_excel_number_only", "Enter a number of 0 or more."),
            )
            validation.add(f"{letter}2:{letter}{LAST_ROW}")
            data.add_data_validation(validation)
        elif key == "barcode":
            # Text, so Excel keeps leading zeros and never shows 13 digits as 6.24E+12.
            for row in range(2, LAST_ROW + 1):
                data[f"{letter}{row}"].number_format = "@"
    data.freeze_panes = "A2"

    instructions.sheet_view.rightToLeft = get_current_language_code() == "ar"
    instructions.column_dimensions["A"].width = 24
    instructions.column_dimensions["B"].width = 12
    instructions.column_dimensions["C"].width = 80
    instructions.column_dimensions["D"].width = 20
    instructions["A1"] = t("opening_template_title", "Opening stock import template")
    instructions["A1"].font = Font(bold=True, size=14)
    instructions.merge_cells("A2:D2")
    instructions["A2"] = t("opening_template_steps")
    instructions["A2"].alignment = WRAP
    instructions.row_dimensions[2].height = 70

    header_row = 4
    headings = (
        t("ui_excel_column", "Column"), t("ui_excel_required", "Required"),
        t("ui_excel_what_to_enter", "What to enter"), t("ui_excel_example", "Example"),
    )
    for index, heading in enumerate(headings, start=1):
        cell = instructions.cell(row=header_row, column=index, value=heading)
        cell.font = Font(bold=True)
        cell.fill = OPTIONAL_FILL
    for row, column in enumerate(guide, start=header_row + 1):
        status = t("ui_excel_required", "Required") if column["required"] else t("ui_excel_optional", "Optional")
        values = (column["label"], status, column["description"], column["example"])
        for index, value in enumerate(values, start=1):
            cell = instructions.cell(row=row, column=index, value=value)
            cell.alignment = WRAP
            if column["required"] and index <= 2:
                cell.font = Font(bold=True)

    row = header_row + len(guide) + 2
    instructions.cell(row=row, column=1, value=t("opening_template_rules", "Rules")).font = Font(bold=True)
    for key in ("opening_template_rule_variants", "opening_template_rule_existing", "opening_template_rule_limits"):
        row += 1
        instructions.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        cell = instructions.cell(row=row, column=1, value=t(key))
        cell.alignment = WRAP
        instructions.row_dimensions[row].height = 32

    workbook.active = 0
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
