"""Validate and stage Opening Stock rows imported from an XLSX workbook."""
from decimal import Decimal
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from django.core.exceptions import ValidationError

from .forms import OpeningStockLineForm
from .models import Category, Product


MAX_IMPORT_BYTES = 2 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 25 * 1024 * 1024
MAX_IMPORT_ROWS = 1000

IMPORT_COLUMNS = (
    ("name", "Name", True),
    ("category", "Category", False),
    ("unit", "Unit", False),
    ("barcode", "Barcode", False),
    ("color", "Color", False),
    ("size", "Size / Spec", False),
    ("cost_usd", "Cost (USD)", False),
    ("markup_percent", "Markup %", False),
    ("price_usd", "Price (USD)", False),
    ("price_lyd_override", "Manual LYD Price", False),
    ("quantity", "Quantity in Storage", True),
)

_HEADER_ALIASES = {
    "name": "name",
    "product": "name",
    "product name": "name",
    "item": "name",
    "item name": "name",
    "category": "category",
    "unit": "unit",
    "barcode": "barcode",
    "color": "color",
    "colour": "color",
    "size": "size",
    "size / spec": "size",
    "size/spec": "size",
    "spec": "size",
    "cost": "cost_usd",
    "cost usd": "cost_usd",
    "cost (usd)": "cost_usd",
    "import cost": "cost_usd",
    "import cost (usd)": "cost_usd",
    "markup": "markup_percent",
    "markup %": "markup_percent",
    "markup percent": "markup_percent",
    "price": "price_usd",
    "price usd": "price_usd",
    "price (usd)": "price_usd",
    "selling price": "price_usd",
    "selling price (usd)": "price_usd",
    "manual lyd": "price_lyd_override",
    "manual lyd price": "price_lyd_override",
    "price lyd": "price_lyd_override",
    "price (lyd)": "price_lyd_override",
    "quantity": "quantity",
    "qty": "quantity",
    "quantity in storage": "quantity",
    "in storage": "quantity",
}


class OpeningStockImportError(ValidationError):
    pass


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _choice(value, choices, field_name):
    raw = _text(value)
    if not raw:
        return ""
    normalized = raw.casefold()
    for code, label in choices:
        if normalized in {str(code).casefold(), str(label).casefold()}:
            return code
    valid = ", ".join(str(code) for code, _label in choices)
    raise OpeningStockImportError(f"Unknown {field_name} '{raw}'. Use one of: {valid}.")


def _category(value):
    raw = _text(value)
    if not raw:
        return None
    matches = list(Category.objects.filter(name__iexact=raw).order_by("pk")[:2])
    if not matches:
        raise OpeningStockImportError(
            f"Category '{raw}' does not exist. Create it first or leave Category blank."
        )
    if len(matches) > 1:
        raise OpeningStockImportError(f"Category '{raw}' is ambiguous.")
    return matches[0]


def _existing_product(name, barcode):
    name_matches = list(Product.objects.filter(name__iexact=name).order_by("pk")[:2])
    barcode_matches = list(Product.objects.filter(barcode=barcode).order_by("pk")[:2]) if barcode else []
    if len(name_matches) > 1:
        raise OpeningStockImportError(f"More than one existing product is named '{name}'.")
    if len(barcode_matches) > 1:
        raise OpeningStockImportError(f"Barcode {barcode} belongs to more than one existing product.")
    by_name = name_matches[0] if name_matches else None
    by_barcode = barcode_matches[0] if barcode_matches else None
    if by_name and by_barcode and by_name.pk != by_barcode.pk:
        raise OpeningStockImportError(
            f"Name matches '{by_name.name}', but barcode belongs to '{by_barcode.name}'."
        )
    product = by_barcode or by_name
    if product and barcode and product.barcode and product.barcode != barcode:
        raise OpeningStockImportError(
            f"'{name}' already uses barcode {product.barcode}; the sheet has {barcode}."
        )
    if product and by_barcode and product.name.casefold() != name.casefold():
        raise OpeningStockImportError(
            f"Barcode {barcode} already belongs to '{product.name}', not '{name}'."
        )
    return product


def _form_errors(form):
    messages = []
    for field, errors in form.errors.items():
        label = form.fields[field].label if field in form.fields else field
        messages.extend(f"{label}: {error}" for error in errors)
    return messages


def _changes(product, cleaned):
    if not product:
        return []
    comparisons = (
        ("Category", product.category_id, getattr(cleaned.get("category"), "pk", None)),
        ("Unit", product.unit, cleaned.get("unit")),
        ("Barcode", product.barcode or "", cleaned.get("barcode") or ""),
        ("Cost (USD)", product.cost_usd, cleaned.get("cost_usd") or Decimal("0")),
        ("Markup %", product.markup_percent, cleaned.get("markup_percent") or Decimal("0")),
        ("Price (USD)", product.price_usd, cleaned.get("price_usd") or Decimal("0")),
        ("Manual LYD", product.price_lyd_override, cleaned.get("price_lyd_override")),
    )
    return [label for label, old, new in comparisons if old != new]


def _primitive(cleaned):
    product_value = cleaned.get("product")
    return {
        "product": getattr(product_value, "pk", product_value) or "",
        "name": cleaned.get("name") or "",
        "category": cleaned["category"].pk if cleaned.get("category") else "",
        "unit": cleaned.get("unit") or Product.UNIT_PIECE,
        "barcode": cleaned.get("barcode") or "",
        "color": cleaned.get("color") or "",
        "size": cleaned.get("size") or "",
        "cost_usd": str(cleaned.get("cost_usd") or Decimal("0")),
        "markup_percent": str(cleaned.get("markup_percent") or Decimal("0")),
        "price_usd": str(cleaned.get("price_usd") or Decimal("0")),
        "price_lyd_override": (
            str(cleaned["price_lyd_override"])
            if cleaned.get("price_lyd_override") is not None else ""
        ),
        "quantity": str(cleaned.get("quantity") or Decimal("0")),
    }


def validate_workbook(upload):
    """Return a preview and form-ready primitive rows; never mutates inventory."""
    if not upload:
        raise OpeningStockImportError("Choose an XLSX file.")
    if upload.size > MAX_IMPORT_BYTES:
        raise OpeningStockImportError("The workbook is larger than 2 MB.")
    if not upload.name.lower().endswith(".xlsx"):
        raise OpeningStockImportError("Only .xlsx workbooks are supported.")

    try:
        from openpyxl import load_workbook

        content = upload.read()
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise OpeningStockImportError("The workbook expands beyond the 25 MB safety limit.")
        workbook = load_workbook(
            BytesIO(content), read_only=True, data_only=True, keep_links=False
        )
        sheet = workbook.active
    except OpeningStockImportError:
        raise
    except BadZipFile as exc:
        raise OpeningStockImportError("The selected file is not a readable XLSX workbook.") from exc
    except Exception as exc:
        raise OpeningStockImportError("The selected file is not a readable XLSX workbook.") from exc

    raw_headers = [_text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1), ())]
    headers = []
    seen_headers = set()
    for raw in raw_headers:
        key = _HEADER_ALIASES.get(raw.casefold())
        headers.append(key)
        if key:
            if key in seen_headers:
                raise OpeningStockImportError(f"The '{raw}' column appears more than once.")
            seen_headers.add(key)
    missing = [label for key, label, required in IMPORT_COLUMNS if required and key not in seen_headers]
    if missing:
        raise OpeningStockImportError("Missing required columns: " + ", ".join(missing) + ".")

    preview = []
    valid_rows = []
    seen_variant_keys = set()
    product_values = {}
    for index, cells in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if index > MAX_IMPORT_ROWS + 1:
            raise OpeningStockImportError(f"A workbook may contain at most {MAX_IMPORT_ROWS} data rows.")
        source = {key: value for key, value in zip(headers, cells) if key}
        if not any(_text(value) for value in source.values()):
            continue

        errors = []
        name = _text(source.get("name"))
        barcode = _text(source.get("barcode"))
        try:
            category = _category(source.get("category"))
            unit = _choice(source.get("unit"), Product.UNIT_CHOICES, "unit") or Product.UNIT_PIECE
            color = _choice(source.get("color"), Product.COLOR_CHOICES, "color")
            product = _existing_product(name, barcode) if name else None
        except OpeningStockImportError as exc:
            category = None
            unit = Product.UNIT_PIECE
            color = ""
            product = None
            errors.extend(exc.messages)

        form_data = {
            "product": product.pk if product else "",
            "name": name,
            "category": category.pk if category else "",
            "unit": unit,
            "barcode": barcode,
            "color": color,
            "size": _text(source.get("size")),
            "cost_usd": _text(source.get("cost_usd")) or "0",
            "markup_percent": _text(source.get("markup_percent")) or "0",
            "price_usd": _text(source.get("price_usd")) or "0",
            "price_lyd_override": _text(source.get("price_lyd_override")),
            "quantity": _text(source.get("quantity")),
        }
        form = OpeningStockLineForm(form_data)
        if not form.is_valid():
            errors.extend(_form_errors(form))

        cleaned = form.cleaned_data if form.is_valid() else None
        if cleaned:
            variant_key = (
                product.pk if product else name.casefold(),
                cleaned.get("color") or "",
                (cleaned.get("size") or "").casefold(),
            )
            if variant_key in seen_variant_keys:
                errors.append("This item/color/size row is duplicated in the workbook.")
            seen_variant_keys.add(variant_key)

            product_key = product.pk if product else name.casefold()
            identity = tuple(
                str(value or "")
                for value in (
                    cleaned.get("category"), cleaned.get("unit"), cleaned.get("barcode"),
                    cleaned.get("cost_usd"), cleaned.get("markup_percent"),
                    cleaned.get("price_usd"), cleaned.get("price_lyd_override"),
                )
            )
            prior = product_values.setdefault(product_key, identity)
            if prior != identity:
                errors.append("Rows for the same item use conflicting product details or prices.")

        changes = _changes(product, cleaned) if cleaned else []
        preview.append({
            "row": index,
            "name": name or "—",
            "variant": " / ".join(filter(None, [color, _text(source.get("size"))])) or "—",
            "quantity": _text(source.get("quantity")) or "—",
            "action": "update" if product else "create",
            "changes": changes,
            "errors": errors,
        })
        if cleaned and not errors:
            valid_rows.append(_primitive(cleaned))

    if not preview:
        raise OpeningStockImportError("The workbook contains no data rows.")
    blocking = sum(bool(row["errors"]) for row in preview)
    return {
        "rows": preview,
        "valid_rows": valid_rows,
        "summary": {
            "total": len(preview),
            "create": sum(row["action"] == "create" and not row["errors"] for row in preview),
            "update": sum(row["action"] == "update" and not row["errors"] for row in preview),
            "conflicts": blocking,
        },
        "can_finalize": blocking == 0 and bool(valid_rows),
    }


def restore_rows(rows):
    """Revalidate staged primitives immediately before the transaction writes."""
    cleaned_rows = []
    for row in rows:
        form = OpeningStockLineForm(row)
        if not form.is_valid():
            raise OpeningStockImportError("The staged import is no longer valid. Preview it again.")
        cleaned_rows.append(form.cleaned_data)
    return cleaned_rows
